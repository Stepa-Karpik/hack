"""Fictional filled local evaluation variants, created before first model scoring."""
from pathlib import Path
import json,hashlib
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
BASE=Path('work/official-expansion-v1')
OUT=BASE/'round3-filled'

def save(doc,name,source,kind,fields):
    doc.add_paragraph('УЧЕБНЫЙ ПРИМЕР. ФИО, ОРГАНИЗАЦИЯ И ДАТЫ ВЫМЫШЛЕНЫ. БЕЗ ПОДПИСИ.')
    path=OUT/(name+'.docx');doc.save(path)
    return {'id':name,'source_id':source,'type':kind,'docx':str(path),'docx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'fields':fields,'unmarked':True}

def main():
    OUT.mkdir(parents=True,exist_ok=True);rows=[]
    for n,(name,group,reason,date) in enumerate([
        ('КУЗНЕЦОВА ОЛЬГА ПЕТРОВНА','Ф-31','семейным обстоятельствам','11.10.2026'),
        ('Алиева-Орлова Тамара Рустамовна','Ф-42','состоянию здоровья','12.10.2026')],1):
        doc=Document(BASE/'originals/isu_leave_2020.docx')
        doc.paragraphs[4].text='\tГруппы № '+group
        doc.paragraphs[5].text='\t'+name+'\n\t(Ф.И.О.)'
        doc.paragraphs[14].text='Прошу предоставить мне академический отпуск по '+reason+'.'
        doc.paragraphs[15].text='';doc.paragraphs[16].text=''
        doc.paragraphs[23].text='Дата: '+date
        fields={'student_name':name,'group':group,'course':None,'reason':reason,'start_date':None,'end_date':None,'application_date':date}
        rows.append(save(doc,'isu_leave_filled_'+str(n),'isu_leave_2020','leave',fields))
    for n,(name,org,grade,date,supervisor) in enumerate([
        ('РОМАНОВ МАКСИМ ВЛАДИМИРОВИЧ','ООО «Правовая поддержка региона»','отлично','18.08.2026','Липатова Инна Борисовна'),
        ('Николаева Мария Константиновна','АО «Объединённый центр правовых решений»','хорошо','19.08.2026','Акимов Константин Андреевич')],1):
        doc=Document(BASE/'originals/kfu_practice_character.docx')
        doc.paragraphs[5].text='кафедра гражданского права'
        doc.paragraphs[10].text='Студента(ки) 4 курса ЮР-41 группы '+name
        doc.paragraphs[11].text='Юридического факультета КФУ очной формы обучения по направлению подготовки 40.03.01 «Юриспруденция» (квалификация «Бакалавр»).'
        p=doc.paragraphs[13]
        p.text=p.text.replace('Ф.И.О. (полностью)',name).replace('с «_»     2020 года по «_»        2020 года','с 01.07.2026 по 12.08.2026').replace('____________________________________________', 'в '+org).replace('___________',org)
        for p in doc.paragraphs:
            if 'Ф.И.О.' in p.text:
                p.text=p.text.replace('Ф.И.О.',name)
        doc.paragraphs[18].text=doc.paragraphs[18].text.replace('____________________',grade)
        doc.paragraphs[21].text='доцент\t\t'+supervisor
        doc.paragraphs[25].text='Дата: '+date
        fields={'student_name':name,'organization':org,'practice_period':'01.07.2026 — 12.08.2026','supervisor':supervisor,'grade':grade,'document_number':None,'issue_date':date}
        rows.append(save(doc,'kfu_practice_filled_'+str(n),'kfu_practice_character','practice',fields))
    report={'scope':'Four fictional filled original-layout derivatives for first-pass evaluation. Names/organizations/dates invented. No signatures/stamps, no model inference or training during creation. PDF/visual QA and immutable input hashes required before scoring. These do not measure real filled student document accuracy.','rows':rows}
    (ROOT/'artifacts/official_round3_filled_manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Four fictional variants prepared; not yet scored')

if __name__=='__main__':main()
