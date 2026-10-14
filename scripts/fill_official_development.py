from pathlib import Path
from docx import Document
from docx.shared import Pt,RGBColor
import hashlib,json

BASE=Path('work/official-expansion-v1')
OUT=BASE/'filled-development';OUT.mkdir(exist_ok=True)
PROJECT=Path(__file__).resolve().parents[1]

def replace(p,text):
    # Preserve paragraph/layout settings and original first-run font.
    if p.runs:
        p.runs[0].text=text
        for r in p.runs[1:]:r.text=''
    else:p.add_run(text)

def cell(d,table,row,col,text):replace(d.tables[table].cell(row,col).paragraphs[0],text)

def save(d,key,source,fields):
    for section in d.sections:
        p=section.footer.paragraphs[0];p.text='УЧЕБНЫЙ ПРИМЕР. ФИО И ДАТЫ ВЫМЫШЛЕНЫ. БЕЗ ПОДПИСИ.'
        for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(100,100,100)
    path=OUT/(key+'.docx');d.save(path)
    return {'id':key,'source_id':source,'institution':source.split('_')[0],'split':'development','type':'leave' if source.startswith('hse') else 'practice','docx':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'fields':fields,'signature':False,'stamp':False,'scope':'Fictional filled variant of official public template; original layout reused locally, no training; not an authentic student document or real camera capture.'}

records=[]
for i,(name,course,reason) in enumerate([('Соколова Надежда Викторовна','3','необходимостью ухода за близким родственником'),('Лаврентьев Денис Павлович','2','изменением семейных обстоятельств и переездом')]):
    source='hse_leave_family';d=Document(BASE/'originals'/(source+'.docx'))
    cell(d,0,2,4,course);cell(d,0,8,2,name)
    replace(d.paragraphs[5],'с 15.10.2026 по 14.10.2027 в связи с '+reason+'.')
    for n in [6,7]:replace(d.paragraphs[n],'')
    cell(d,1,2,2,'10.10.2026')
    records.append(save(d,f'hse_family_filled_{i+1}',source,{'student_name':name,'group':None,'course':course,'reason':'по семейным обстоятельствам в связи с '+reason,'start_date':'15.10.2026','end_date':'14.10.2027','application_date':'10.10.2026'}))

for i,(name,org,grade,supervisor) in enumerate([('Селезнёва Вера Константиновна','ООО «Технологии северного города»','отлично','Гордеев Николай Алексеевич'),('Фролов Кирилл Евгеньевич','АО «Центр инженерных решений»','хорошо','Матвеева Анна Сергеевна')]):
    source='rudn_practice_character';d=Document(BASE/'originals'/(source+'.docx'))
    replace(d.paragraphs[4],name)
    period='15.06.2026 г. по 12.07.2026 г.' if i==0 else '15 июня 2026 г. по 12 июля 2026 г.'
    replace(d.paragraphs[6],'в период '+period+' проходил(а)')
    replace(d.paragraphs[7],'производственную практику в '+org+'.')
    replace(d.paragraphs[10],'')
    replace(d.paragraphs[12],'При прохождении практики выполнял(а) задания по анализу данных и подготовке отчётных материалов.')
    for n in [13,14]:replace(d.paragraphs[n],'')
    replace(d.paragraphs[17],'Обучающийся '+name)
    replace(d.paragraphs[19],'по результатам прохождения практики заслуживает оценки «'+grade+'»')
    replace(d.paragraphs[23],'от организации (предприятия) '+supervisor+', руководитель отдела')
    replace(d.paragraphs[27],'09.08.2026')
    records.append(save(d,f'rudn_character_filled_{i+1}',source,{'student_name':name,'organization':org,'practice_period':'15.06.2026 - 12.07.2026','supervisor':supervisor,'grade':grade,'document_number':None,'issue_date':'09.08.2026'}))

report={'scope':'Four privately filled official template variants, two source institutions already designated development. Values locked before OCR. No source logos/signatures/stamps added. Source reuse permissions not verified; keep originals/derivatives outside release.', 'records':records}
(PROJECT/'artifacts/official_filled_truth.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('Four fictional filled DOCX written; truth frozen before inference')
