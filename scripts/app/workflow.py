"""Application workflow helpers: submit warning text, name check, Excel export, PDF protocol, history labels."""
import io
import re
import time
from datetime import datetime
from pathlib import Path

FONTS = Path(__file__).resolve().parent / 'fonts'

TYPE_NAMES = {'study_certificate': 'Справка об обучении', 'academic_leave_application': 'Заявление на академический отпуск',
              'internship_certificate': 'Справка о практике'}
STATE_NAMES = {'draft': 'Черновик', 'submitted': 'Отправлен в деканат', 'review': 'На рассмотрении',
               'approved': 'Одобрен', 'returned': 'Возвращён на доработку', 'closed': 'Заявка закрыта'}
EVENT_NAMES = {'uploaded': 'Документ загружен', 'processed': 'Документ проверен автоматически', 'submitted': 'Отправлен в деканат',
               'review': 'Взят на рассмотрение', 'approved': 'Одобрен', 'returned': 'Возвращён на доработку',
               'reset': 'Решение отменено', 'replaced': 'Файл заменён', 'edited': 'Изменён в редакторе', 'closed': 'Заявка закрыта',
               'type_changed': 'Изменён тип документа', 'auto_approved': 'Одобрен автоматически'}

# «не хватает …» — every field in the genitive, so the sentence reads naturally
GENITIVE = {'student_name': 'ФИО студента', 'university': 'названия вуза', 'program': 'направления подготовки',
            'study_form': 'формы обучения', 'study_start': 'даты начала обучения', 'study_end': 'даты окончания обучения',
            'document_number': 'номера справки', 'issue_date': 'даты выдачи', 'group': 'учебной группы', 'course': 'курса',
            'leave_reason': 'причины отпуска', 'leave_start': 'даты начала отпуска', 'leave_end': 'даты окончания отпуска',
            'application_date': 'даты заявления', 'organization': 'названия организации',
            'internship_start': 'даты начала практики', 'internship_end': 'даты окончания практики',
            'supervisor_name': 'ФИО руководителя практики', 'grade': 'оценки'}
MARK_GENITIVE = {('signature_present', 'study_certificate'): 'подписи сотрудника вуза',
                 ('signature_present', 'academic_leave_application'): 'вашей подписи',
                 ('signature_present', 'internship_certificate'): 'подписи руководителя практики',
                 ('seal_present', 'study_certificate'): 'печати вуза',
                 ('seal_present', 'internship_certificate'): 'печати организации'}


def join_ru(items):
    items = list(items)
    if len(items) <= 1:
        return ''.join(items)
    return ', '.join(items[:-1]) + ' и ' + items[-1]


def field_genitive(key, kind, labels=None):
    if (key, kind) in MARK_GENITIVE:
        return MARK_GENITIVE[(key, kind)]
    if key == 'signature_present':
        return 'подписи'
    if key == 'seal_present':
        return 'печати'
    return GENITIVE.get(key) or f'поля «{(labels or {}).get(key, key)}»'


def submit_warning(analysis, name_check=None, labels=None):
    """A human sentence for the student before sending an incomplete document, or None."""
    kind = analysis.get('document_type')
    parts = []
    validation = analysis.get('validation') or {}
    missing = validation.get('missing_fields') or []
    quality = analysis.get('quality') or {}
    if quality.get('readable') is False:
        parts.append('Снимок читается плохо: часть сведений могла не распознаться, хотя в документе она есть.')
    if not analysis.get('type_matches', True):
        parts.append('Мы не узнали в файле выбранный тип документа, поэтому сведения могли не найтись.')
    if missing:
        rules = [s.split(':', 1)[1] for s in validation.get('sources') or []
                 if s.startswith('rule:') and re.fullmatch(r'\d+(\.\d+)*', s.split(':', 1)[1])]
        what = join_ru(field_genitive(k, kind, labels) for k in missing)
        basis = (' — этого требует пункт ' + rules[0] if len(rules) == 1 else
                 ' — этого требуют пункты ' + join_ru(rules) if rules else '')
        parts.append(f'Мы проверили документ: в нём не хватает {what}{basis} Положения.')
    if name_check and name_check.get('status') == 'mismatch':
        parts.append(f'ФИО в документе («{name_check["document"]}») не совпадает с ФИО в вашем профиле '
                     f'(«{name_check["profile"]}»).')
    if not parts:
        return None
    return ' '.join(parts) + (' С таким документом заявление, скорее всего, вернут на доработку. Лучше добавить недостающее '
                              'и заменить файл. Если вы уверены, что всё на месте и система ошиблась, всё равно отправьте '
                              'документ — сотрудник деканата проверит его вручную.')


def _name_tokens(value):
    value = (value or '').casefold().replace('ё', 'е')
    return [t for t in re.findall(r'[а-яa-z]+', value) if len(t) > 1]


def name_check(profile_name, document_name):
    """Compare the account ФИО with the student name read from the document (surname + first name)."""
    profile, document = _name_tokens(profile_name), _name_tokens(document_name)
    if len(profile) < 2 or not document or document_name in (None, 'не найдено'):
        return None
    # declension-tolerant: compare word stems (first 3 letters) of surname and first name, in any order
    stems = {t[:3] for t in document}
    matched = sum(1 for t in profile[:2] if t[:3] in stems)
    return {'status': 'ok' if matched == 2 else 'mismatch', 'profile': profile_name, 'document': document_name}


def fmt_time(ts):
    return datetime.fromtimestamp(ts).strftime('%d.%m.%Y %H:%M') if ts else ''


def who(row, prefix='user'):
    if not row.get(f'{prefix}_email'):
        return 'система'
    return (row.get(f'{prefix}_name') or row[f'{prefix}_email']) + f' ({row[f"{prefix}_email"]})'


# ---------------------------------------------------------------- Excel
def export_xlsx(documents, events, type_names):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = 'Заявки'
    head = ['Заявка', 'Студент', 'Почта', 'Группа', 'Тип документа', 'Файл', 'Найдено', 'Комплектность', 'Статус',
            'Отправлено', 'Попыток', 'Решение принял', 'Дата решения', 'Комментарий деканата', 'Комментарий студента']
    ws.append(head)
    for d in documents:
        ws.append([d['slug'], d.get('owner_name') or '', d.get('owner_email') or '', d.get('owner_group') or '',
                   type_names.get(d.get('document_type') or d['requested_type'], d.get('document_type') or ''),
                   d['filename'], f"{d.get('found') or 0}/{d.get('total') or 0}",
                   'полный' if d.get('valid') else 'есть пропуски', STATE_NAMES.get(d.get('state'), d.get('state')),
                   fmt_time(d.get('submitted_at')), d.get('attempts') or 0, d.get('decider') or '',
                   fmt_time(d.get('approved_at') or d.get('decided_at')), d.get('decision_comment') or '',
                   d.get('student_comment') or ''])
    log = wb.create_sheet('Журнал')
    log.append(['Время', 'Заявка', 'Событие', 'Кто', 'Роль', 'Подробности'])
    for e in events:
        detail = e.get('detail') or {}
        text = '; '.join(f'{k}: {v}' for k, v in detail.items() if v not in (None, '', []))
        log.append([fmt_time(e['at']), e['slug'], EVENT_NAMES.get(e['kind'], e['kind']), who(e),
                    {'student': 'студент', 'staff': 'деканат', 'admin': 'администратор'}.get(e.get('user_role'), ''), text])
    for sheet in (ws, log):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='6D5BD0')
            cell.alignment = Alignment(vertical='center')
        sheet.freeze_panes = 'A2'
        for column in sheet.columns:
            width = max(len(str(c.value or '')) for c in column[:200])
            sheet.column_dimensions[column[0].column_letter].width = min(48, max(10, width + 2))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- PDF protocol
def protocol_pdf(doc, analysis, events, labels, type_names, preview=None, display=None):
    from fpdf import FPDF
    pdf = FPDF(format='A4')
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_font('DejaVu', '', str(FONTS / 'DejaVuSans.ttf'))
    pdf.add_font('DejaVu', 'B', str(FONTS / 'DejaVuSans-Bold.ttf'))
    pdf.add_page()
    pdf.set_font('DejaVu', 'B', 15)
    pdf.cell(0, 9, 'Протокол предварительной проверки документа', new_x='LMARGIN', new_y='NEXT')
    pdf.set_font('DejaVu', '', 9)
    pdf.set_text_color(110, 110, 110)
    pdf.cell(0, 5, f'Онлайн-деканат · сформирован {fmt_time(time.time())} · заявка {doc["slug"]}', new_x='LMARGIN', new_y='NEXT')
    pdf.set_text_color(0, 0, 0)
    pdf.ln(3)

    def row(key, value, bold=False):
        pdf.set_font('DejaVu', '', 10)
        pdf.set_text_color(90, 90, 90)
        pdf.cell(55, 6.5, key)
        pdf.set_text_color(0, 0, 0)
        pdf.set_font('DejaVu', 'B' if bold else '', 10)
        pdf.multi_cell(0, 6.5, str(value), new_x='LMARGIN', new_y='NEXT')

    kind = analysis.get('document_type') if analysis else doc.get('requested_type')
    row('Студент', f"{doc.get('owner_name') or ''} {('<' + doc['owner_email'] + '>') if doc.get('owner_email') else ''}".strip())
    row('Документ', f"{type_names.get(kind, kind)} · {doc['filename']} · версия {doc.get('version') or 1}")
    row('Статус заявки', STATE_NAMES.get(doc.get('state'), doc.get('state')), bold=True)
    if doc.get('decision_comment'):
        row('Комментарий деканата', doc['decision_comment'])
    if analysis:
        validation = analysis['validation']
        missing = validation.get('missing_fields') or []
        result = ('Все обязательные сведения найдены' if validation['status'] == 'document_ok'
                  else 'Не хватает: ' + ', '.join(labels.get(k, k) for k in missing))
        rules = ', '.join(s.split(':', 1)[1] for s in validation.get('sources') or [])
        row('Предварительная проверка', f'{result} (п. {rules})' if rules else result, bold=True)
        pdf.ln(2)
        pdf.set_font('DejaVu', 'B', 11)
        pdf.cell(0, 7, 'Сведения из документа', new_x='LMARGIN', new_y='NEXT')
        verified = doc.get('verified') or {}
        for key, value in analysis['fields'].items():
            shown = (display or {}).get(key) or value
            if value is True:
                shown = 'обнаружена'
            elif value is False:
                shown = 'не обнаружена'
            note = ' · сверено вручную' if verified.get(key) else (' · проверьте' if (analysis.get('review') or {}).get(key) else '')
            row(labels.get(key, key), f'{shown}{note}')
    pdf.ln(2)
    pdf.set_font('DejaVu', 'B', 11)
    pdf.cell(0, 7, 'История заявки', new_x='LMARGIN', new_y='NEXT')
    for e in events:
        detail = e.get('detail') or {}
        extra = detail.get('comment') or detail.get('filename') or ''
        pdf.set_font('DejaVu', '', 9)
        pdf.multi_cell(0, 5.5, f"{fmt_time(e['at'])}  {EVENT_NAMES.get(e['kind'], e['kind'])} — {who(e)}"
                       + (f': {extra}' if extra else ''), new_x='LMARGIN', new_y='NEXT')
    if preview and Path(preview).exists():
        try:
            from PIL import Image
            with Image.open(preview) as check:
                check.verify()
            pdf.add_page()
            pdf.set_font('DejaVu', 'B', 11)
            pdf.cell(0, 7, 'Изображение документа', new_x='LMARGIN', new_y='NEXT')
            pdf.image(str(preview), x=15, w=180)
        except Exception:  # the protocol is still useful without the picture
            pass
    pdf.set_y(-20)
    pdf.set_font('DejaVu', '', 8)
    pdf.set_text_color(120, 120, 120)
    pdf.multi_cell(0, 4, 'Результат автоматической проверки носит предварительный характер. Окончательное решение принимает '
                         'сотрудник деканата (п. 1.3 Положения).')
    return bytes(pdf.output())
